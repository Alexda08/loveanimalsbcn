# -*- coding: utf-8 -*-
"""Los adoptados del Excel nuevo de Carla (04-10-2026).

Catorce filas llevan «ADOPTADO/A» pegado al nombre. Se hace con ellas lo que
hay que hacer para que salgan en el muro de finales felices:

  * las que ya están en la tienda  -> estado «adoptado» (y se publican si estaban
    en borrador, que Liquid no ve los borradores);
  * Simba «abuelo» y Fabrizzio, que no estaban (el Excel viejo decía NO PONERLO)
    -> se crean con los datos de su fila y sus fotos;
  * Behia, que estaba en borrador y sin ninguna foto -> con la foto del zip nuevo;
  * las que no estaban en ningún álbum se meten en los que dice el Excel, que
    sin álbum la ficha no sabe de dónde viene.

No se inventa ninguna fecha de adopción: Carla no la ha dado, y el sello del
muro sabe quedarse en «Adoptado» a secas.

    python adoptados.py mirar
    python adoptados.py aplicar
"""
import os, sys, io, json, re, time, zipfile, unicodedata, requests

API = "https://loveanimalsbcn.myshopify.com/admin/api/2026-01/graphql.json"
TOKEN = os.environ["LA_TOKEN_REAL"]
DOCS = r"C:\Users\desca\OneDrive\Desktop\loveanimalsbcn\docs\img_loveanimalsbcn"
ZIP_NUEVO = r"C:\Users\desca\Downloads\Simba, Fabrizzio, Behia .zip"
XLSX = r"C:\Users\desca\Downloads\Love Animals BCN - fichas de los animales.xlsx"

# fila del Excel -> handle en la tienda (existentes o por crear)
FILAS = {5: "simba-abuelo", 58: "fabrizzio", 99: "behia", 169: "saitama", 172: "seimus",
         180: "thorin", 181: "gnar", 183: "kuka-kuqui", 189: "pepinillos", 194: "ninu-siames",
         201: "felipe", 209: "laks", 213: "bulma-cachorra", 218: "bibi-sepi"}
NUEVOS = ("simba-abuelo", "fabrizzio")

# fotos de los animales que no las tienen aún: (zip, nombre dentro del zip)
FOTOS = {
    "simba-abuelo": [(os.path.join(DOCS, "FOTOS PERROS ABUELOS WEB.zip"), "AF85D609-5F5C-4BCA-BA49-13F846EABE8E.png"),
                     (os.path.join(DOCS, "FOTOS PERROS ABUELOS WEB.zip"), "4679C08A-8DF2-4260-B051-B807F906395A.png"),
                     (os.path.join(DOCS, "FOTOS PERROS ABUELOS WEB.zip"), "ba1f5760-eb2f-496c-9df0-738f9c9a2bfa.jpeg")],
    # el zip nuevo trae una foto por cabeza: la de la mujer con el pitbull del pañuelo es
    # el PPP (Fabrizzio) y la del perro grande blanco y negro es la «vaquita» mastín (Behia)
    "fabrizzio": [(ZIP_NUEVO, "408c0ce6-8174-4962-b371-41ae8b01fdde.jpeg")],
    "behia": [(ZIP_NUEVO, "IMG_2913.jpeg")],
}

ALBUM_HANDLE = {"ABUELOS": "abuelos", "DESDE 2022 ESPERANDO FAMILIA": "desde-2022-esperando-familia",
                "DESDE 2024 ESPERANDO FAMILIA": "desde-2024-esperando-familia",
                "PPP ADULTOS": "ppp-adultos", "MESTIZOS": "mestizos", "GATOS": "gatos"}
MEMBRESIA = {  # lo que dice el Excel para los que no estaban en ningún álbum
    "simba-abuelo": ["ABUELOS", "DESDE 2022 ESPERANDO FAMILIA"],
    "fabrizzio": ["DESDE 2022 ESPERANDO FAMILIA", "PPP ADULTOS"],
    "behia": ["DESDE 2024 ESPERANDO FAMILIA", "MESTIZOS"],
    "saitama": ["GATOS"], "thorin": ["GATOS"],
}
ARREGLOS_TEXTO = {u"Simba fue rescatad de la calle": u"Simba fue rescatado de la calle"}


def gql(q, v=None):
    for n in range(6):
        r = requests.post(API, json={"query": q, "variables": v or {}},
                          headers={"X-Shopify-Access-Token": TOKEN}, timeout=180)
        if r.status_code in (429, 500, 502, 503):
            time.sleep(2 * (n + 1)); continue
        d = r.json()
        if any(e.get("extensions", {}).get("code") == "THROTTLED" for e in d.get("errors") or []):
            time.sleep(3); continue
        if d.get("errors"):
            raise SystemExit("GraphQL: " + json.dumps(d["errors"], ensure_ascii=False)[:700])
        return d["data"]
    raise SystemExit("la API no responde")


def errs(r, que):
    if r["userErrors"]:
        raise SystemExit("%s: %s" % (que, json.dumps(r["userErrors"], ensure_ascii=False)))


def tienda():
    out, cur = {}, None
    while True:
        d = gql("""query($c:String){ metaobjects(type:"animal", first:100, after:$c){
            nodes{ id handle capabilities{ publishable{ status } } fields{ key value } }
            pageInfo{ hasNextPage endCursor } } }""", {"c": cur})["metaobjects"]
        for m in d["nodes"]:
            out[m["handle"]] = m
        if not d["pageInfo"]["hasNextPage"]:
            return out
        cur = d["pageInfo"]["endCursor"]


def excel():
    import openpyxl
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Animales"]
    cab = [str(c.value or "").strip() for c in ws[1]]
    return {n: {cab[i]: ws[n][i].value for i in range(len(cab)) if cab[i]} for n in FILAS}


def rich(texto):
    parrafos = [p.strip() for p in re.split(r"\n+", texto or "") if p.strip()]
    return json.dumps({"type": "root", "children": [
        {"type": "paragraph", "children": [{"type": "text", "value": p}]} for p in parrafos]}, ensure_ascii=False)


def sino(v):
    return {"sí": "si", "si": "si", "no": "no", "consultar": "consultar"}.get(str(v or "").strip().lower(), "consultar")


def campos_de(h, fila, gids):
    nombre = re.sub(r"\s*ADOPTAD[OA]S?\s*$", "", str(fila["Nombre *"])).strip()
    hist = str(fila.get("Historia") or "")
    for a, b in ARREGLOS_TEXTO.items():
        hist = hist.replace(a, b)
    car = [c.strip() for c in str(fila.get("Carácter") or "").split(",") if c.strip()]
    f = [("nombre", nombre), ("frase", str(fila.get("Frase") or "")),
         ("fecha_entrada", str(fila["Fecha de entrada *"])[:10]), ("edad", str(fila.get("Edad") or "")),
         ("tamano", str(fila.get("Tamaño") or "")), ("sexo", str(fila.get("Sexo") or "").lower()),
         ("ppp", "true" if str(fila.get("PPP") or "").strip().lower() in ("sí", "si") else "false"),
         ("convive_ninos", sino(fila.get("Convive con niños"))),
         ("convive_perros", sino(fila.get("Convive con perros"))),
         ("convive_gatos", sino(fila.get("Convive con gatos"))),
         ("historia", rich(hist)), ("caracter", json.dumps(car, ensure_ascii=False)),
         ("estado", "adoptado"), ("galeria", json.dumps(gids))]
    return [{"key": k, "value": v} for k, v in f]


def sube(h, quedan):
    """sube las fotos de `h` que aún no estén en la biblioteca. Devuelve los gid."""
    entrada, datos = [], []
    for i, (zp, interno) in enumerate(FOTOS[h], len(quedan) + 1):
        with zipfile.ZipFile(zp) as z:
            b = z.read(interno)
        ext = os.path.splitext(interno)[1].lower().replace(".jpeg", ".jpg")
        nombre = "%s-%d%s" % (h, i, ext)
        ya = gql("""query($q:String!){ files(first:2, query:$q){ nodes{ id } } }""",
                 {"q": "filename:" + nombre})["files"]["nodes"]
        if ya:
            quedan.append(ya[0]["id"]); continue
        datos.append((nombre, b))
        entrada.append({"filename": nombre, "mimeType": "image/png" if ext == ".png" else "image/jpeg",
                        "resource": "IMAGE", "httpMethod": "POST", "fileSize": str(len(b))})
    if not entrada:
        return quedan
    d = gql("""mutation($input:[StagedUploadInput!]!){ stagedUploadsCreate(input:$input){
        stagedTargets{ url resourceUrl parameters{ name value } } userErrors{ field message } } }""",
            {"input": entrada})["stagedUploadsCreate"]
    errs(d, "stagedUploads")
    ficheros = []
    for (nombre, b), t in zip(datos, d["stagedTargets"]):
        r = requests.post(t["url"], data=[(p["name"], p["value"]) for p in t["parameters"]],
                          files={"file": (nombre, b)}, timeout=600)
        if r.status_code not in (200, 201, 204):
            raise SystemExit("subida %s: %s" % (nombre, r.status_code))
        ficheros.append({"originalSource": t["resourceUrl"], "contentType": "IMAGE",
                         "alt": u"%s, de Love Animals BCN" % h.replace("-", " ").title()})
    d = gql("""mutation($files:[FileCreateInput!]!){ fileCreate(files:$files){
        files{ id } userErrors{ field message } } }""", {"files": ficheros})["fileCreate"]
    errs(d, "fileCreate")
    nuevos = [f["id"] for f in d["files"]]
    for _ in range(40):
        ns = gql("""query($ids:[ID!]!){ nodes(ids:$ids){ ... on MediaImage { id fileStatus } } }""",
                 {"ids": nuevos})["nodes"]
        if any(n and n["fileStatus"] == "FAILED" for n in ns):
            raise SystemExit("una foto falló al procesarse")
        if all(n and n["fileStatus"] == "READY" for n in ns):
            break
        time.sleep(3)
    return quedan + nuevos


def main():
    hazlo = len(sys.argv) > 1 and sys.argv[1] == "aplicar"
    S, X = tienda(), excel()
    albums = {a["handle"]: a for a in gql("""{ metaobjects(type:"album", first:50){ nodes{ id handle
        fields{ key value } } } }""")["metaobjects"]["nodes"]}
    ids = {}   # handle -> gid, para meter a los nuevos en los álbumes
    for fila, h in FILAS.items():
        cur = S.get(h)
        estado = ({f["key"]: f["value"] for f in cur["fields"]}.get("estado") if cur else "(nuevo)")
        pub = cur["capabilities"]["publishable"]["status"] if cur else "(nuevo)"
        print("fila %-3d %-16s hoy: %-12s %-8s -> adoptado%s%s" % (
            fila, h, estado, pub, "" if pub in ("ACTIVE",) else " + publicar",
            " + foto nueva" if h in FOTOS else ""))
        if not hazlo:
            continue
        if cur is None:
            gids = sube(h, [])
            d = gql("""mutation($m:MetaobjectCreateInput!){ metaobjectCreate(metaobject:$m){
                metaobject{ id handle } userErrors{ field message } } }""",
                    {"m": {"type": "animal", "handle": h, "fields": campos_de(h, X[fila], gids),
                           "capabilities": {"publishable": {"status": "ACTIVE"}}}})["metaobjectCreate"]
            errs(d, "crear " + h)
            ids[h] = d["metaobject"]["id"]
            continue
        campos = [{"key": "estado", "value": "adoptado"}]
        if h in FOTOS:
            actuales = json.loads({f["key"]: f["value"] for f in cur["fields"]}.get("galeria") or "[]")
            if not actuales:
                campos.append({"key": "galeria", "value": json.dumps(sube(h, []))})
        d = gql("""mutation($id:ID!,$m:MetaobjectUpdateInput!){ metaobjectUpdate(id:$id, metaobject:$m){
            metaobject{ handle } userErrors{ field message } } }""",
                {"id": cur["id"], "m": {"fields": campos, "capabilities": {"publishable": {"status": "ACTIVE"}}}}
                )["metaobjectUpdate"]
        errs(d, "actualizar " + h)
        ids[h] = cur["id"]

    # --- álbumes
    for h, titulos in MEMBRESIA.items():
        for t in titulos:
            a = albums[ALBUM_HANDLE[t]]
            c = {f["key"]: f["value"] for f in a["fields"]}
            lista = json.loads(c.get("animales") or "[]")
            cola = json.loads(c.get("animales_mas") or "[]")
            gid = ids.get(h) or (S[h]["id"] if h in S else None)
            if gid is None or gid in lista or gid in cola:
                print("   álbum %-30s %s: ya estaba o no hay id" % (ALBUM_HANDLE[t], h)); continue
            clave, nueva = ("animales", lista + [gid]) if len(lista) < 50 else ("animales_mas", cola + [gid])
            print("   álbum %-30s + %-14s (en %s, %d)" % (ALBUM_HANDLE[t], h, clave, len(nueva)))
            if hazlo:
                d = gql("""mutation($id:ID!,$m:MetaobjectUpdateInput!){ metaobjectUpdate(id:$id, metaobject:$m){
                    metaobject{ handle } userErrors{ field message } } }""",
                        {"id": a["id"], "m": {"fields": [{"key": clave, "value": json.dumps(nueva)}]}})["metaobjectUpdate"]
                errs(d, "álbum " + t)
                c[clave] = json.dumps(nueva)
                a["fields"] = [{"key": k, "value": v} for k, v in c.items()]
    print("\n" + ("APLICADO" if hazlo else "(solo mirando)"))


if __name__ == "__main__":
    main()
