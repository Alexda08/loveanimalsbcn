# -*- coding: utf-8 -*-
"""Cotejo del Excel de Carla contra lo que ha ido diciendo por WhatsApp.

No toca nada: solo dice donde el Excel y sus mensajes no cuadran.
"""
import os, re, unicodedata, collections
import openpyxl

XLSX = os.path.join(r"C:\Users\desca\OneDrive\Desktop\loveanimalsbcn", "docs",
                    "Love Animals BCN - fichas de los animales.xlsx")

COLS = ["n", "nombre", "album", "estado", "entrada", "adopcion", "sexo", "edad", "tamano",
        "ppp", "ninos", "perros", "gatos", "caracter", "frase", "historia", "fotos", "notas", "cod"]

# lo que Carla dijo por WhatsApp el 03-09-2026, tal cual
ADOPTADOS = ["Behia", "Rusty", "Kenia", "Simba abuelo", "Fabrizzio",
             "Saitama", "Dustin", "Bony", "Xulo", "Thorin"]
NO_PUBLICAR = ["Loki", "Negrita", "Ramen", "Liona", "Carmela"]
SIN_FOTO = ["Crowley", "Andres", "Chiquitita", "Kaur", "Katsuki"]
SI_PUBLICAR = ["Kay", "Bibi & Sepi"]
DUPLICADA = ["Sakura"]
FALLECIDOS = ["Muso", "Joker"]
# los que dio por perdidos «del final»
PERDIDOS = ["xip", "kuka", "kuqui", "pitu", "patu", "ninu", "pepinillos", "canica",
            "mix", "max", "hermione", "ron", "harry", "huran", "felipe", "reina",
            "aixa", "flipi", "tracy", "enzo"]


def pelado(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def main():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Animales"]
    filas = []
    for r in range(3, ws.max_row + 1):
        v = {k: ws.cell(r, i + 1).value for i, k in enumerate(COLS)}
        if not str(v["nombre"] or "").strip():
            continue
        v["fila"] = r
        v["_nombre"] = re.sub(r"\s+", " ", str(v["nombre"]).strip())
        filas.append(v)
    print("filas con nombre: %d (de la 3 a la %d)\n" % (len(filas), ws.max_row))

    por_nombre = collections.defaultdict(list)
    for v in filas:
        por_nombre[pelado(v["_nombre"])].append(v)

    def busca(nombre):
        p = pelado(nombre)
        if p in por_nombre:
            return por_nombre[p]
        return [v for k, vs in por_nombre.items() if p in k or k in p for v in vs]

    marcados = [v for v in filas
                if re.search(r"\bno\s+pon(er|erlo|erla|gais|ga)\b", pelado(v["notas"]))]
    print("== 1. FILAS MARCADAS «NO PONERLO» EN EL EXCEL (%d) ==" % len(marcados))
    dicho = {pelado(x) for x in NO_PUBLICAR + DUPLICADA}
    for v in marcados:
        p = pelado(v["_nombre"])
        if any(d in p or p in d for d in dicho):
            estado = "ok, Carla lo confirma"
        elif any(pelado(a) in p or p in pelado(a) for a in ADOPTADOS):
            estado = "ADOPTADO, ya no hace falta la marca"
        elif any(pelado(a) in p or p in pelado(a) for a in SI_PUBLICAR):
            estado = "!! CARLA LO QUIERE PUBLICADO"
        else:
            estado = "?? no lo ha mencionado"
        print("   fila %3d  %-18s  %-34s  %s" % (v["fila"], v["_nombre"][:18], estado,
                                                 re.sub(r"\s+", " ", str(v["notas"]))[:44]))

    print("\n== 2. LO QUE CARLA DIJO, FILA A FILA ==")
    grupos = [("adoptado", ADOPTADOS), ("no publicar", NO_PUBLICAR),
              ("falta foto", SIN_FOTO), ("si publicar", SI_PUBLICAR),
              ("duplicada", DUPLICADA), ("fallecido", FALLECIDOS)]
    for etiqueta, nombres in grupos:
        for n in nombres:
            vs = busca(n)
            if not vs:
                print("   %-12s %-16s -- NO ESTA EN EL EXCEL" % (etiqueta, n))
                continue
            for v in vs:
                marca = "NO PONERLO" if v in marcados else ""
                print("   %-12s %-16s fila %3d | fotos=%-3s | estado=%-12s | %s"
                      % (etiqueta, n, v["fila"], "si" if str(v["fotos"] or "").strip() else "NO",
                         str(v["estado"] or "")[:12], marca))

    print("\n== 3. LOS QUE DIO POR PERDIDOS «DEL FINAL» ==")
    for n in PERDIDOS:
        vs = [v for k, vs2 in por_nombre.items() if n in k for v in vs2]
        if not vs:
            print("   %-12s -- NO ESTA" % n)
        else:
            print("   %-12s fila %s" % (n, ", ".join(str(v["fila"]) for v in vs)))

    print("\n== 4. NOMBRES REPETIDOS (por si hay mas Sakuras) ==")
    hay = False
    for k, vs in sorted(por_nombre.items()):
        if len(vs) > 1:
            hay = True
            print("   %-22s filas %s" % (k, [v["fila"] for v in vs]))
    if not hay:
        print("   ninguno exacto")

    print("\n== 5. ESTADO Y FECHA DE ADOPCION ==")
    est = collections.Counter(pelado(v["estado"]) for v in filas)
    for k, n in est.most_common():
        print("   %-16s %d" % (k or "(vacio)", n))
    con_fecha = [v for v in filas if str(v["adopcion"] or "").strip()]
    print("   con fecha de adopcion:", len(con_fecha))
    for v in con_fecha[:10]:
        print("      fila %3d %-18s %s" % (v["fila"], v["_nombre"], v["adopcion"]))

    print("\n== 6. FILAS SIN FOTO ==")
    for v in filas:
        if not str(v["fotos"] or "").strip():
            marca = "NO PONERLO" if v in marcados else ""
            print("   fila %3d  %-20s  %-30s %s" % (v["fila"], v["_nombre"][:20],
                                                    str(v["album"] or "")[:30], marca))

    print("\n== 7. FILAS SIN FECHA DE ENTRADA ==")
    for v in filas:
        if not str(v["entrada"] or "").strip():
            print("   fila %3d  %-20s" % (v["fila"], v["_nombre"][:20]))


if __name__ == "__main__":
    main()
