# -*- coding: utf-8 -*-
"""Genera el Excel que se le manda a Carla para las fichas de los animales.

Va en blanco: solo la primera fila lleva un ejemplo relleno para que se vea el
formato. La última columna guarda el código interno de la ficha (vacío = ficha
nueva); es lo que permite reimportar sin duplicar.
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding="utf-8")

SALIDA = r"C:\Users\desca\OneDrive\Desktop\Love Animals BCN - fichas de los animales.xlsx"

# paleta del proyecto
VINO = "6E4A44"
CHOCO = "4A3628"
DORADO = "B08D57"
CREMA = "F7F2E8"
NUDE = "EFE8DA"
VERDE = "5F7A5B"
GRIS = "9A9186"

FILAS = 200               # filas para rellenar
ULTIMA = FILAS + 1        # la 1 es la cabecera
EJEMPLO = 2               # fila del ejemplo

ALBUMES = [
    "NECESITAMOS CASA DE ACOGIDA", "ABUELOS", "LOS MÁS VETERANOS",
    "PPP JÓVENES", "PPP ADULTOS", "MESTIZOS", "GATOS",
    "DESDE 2021 ESPERANDO FAMILIA", "DESDE 2022 ESPERANDO FAMILIA",
    "DESDE 2023 ESPERANDO FAMILIA", "DESDE 2024 ESPERANDO FAMILIA",
]

# ---------------------------------------------------------------- columnas

# (título, ancho, obligatoria, ayuda, validación, ejemplo)
COLUMNAS = [
    ("Nº", 5, False, "Solo para contar. No hace falta tocarlo.", None, "ej."),
    ("Nombre *", 16, True,
     "El nombre del animal, tal y como queréis que salga en la web.", None, "Woody"),
    ("Álbum(es)", 30, False,
     "En qué álbumes de la web sale. Varios separados por punto y coma. "
     "Los álbumes válidos están en la hoja «Listas».", None,
     "PPP ADULTOS; DESDE 2021 ESPERANDO FAMILIA"),
    ("Estado *", 14, True,
     "En adopción = busca familia · En acogida = está en una casa de acogida · "
     "Adoptado = ya tiene familia (sale en el muro de finales felices).",
     '"En adopción,En acogida,Adoptado"', "En adopción"),
    ("Fecha de entrada *", 15, True,
     "Cuándo llegó al refugio. Es lo que hace funcionar el contador de días esperando. "
     "Si no recordáis el día exacto, poned el día 1 del mes (ej. 01/03/2021).",
     None, date(2021, 3, 1)),
    ("Fecha de adopción", 15, False,
     "Solo si el estado es «Adoptado». Sirve para ordenar el muro de finales felices.",
     None, None),
    ("Sexo", 10, False, "Macho o hembra.", '"Macho,Hembra"', "Macho"),
    ("Edad", 18, False,
     "Tal cual queráis que se lea: «2 años», «cachorro», «+9 años», «adulto (5-9 años)».",
     None, "6 años"),
    ("Tamaño", 12, False, "Pequeño, mediano o grande.", '"Pequeño,Mediano,Grande"', "Grande"),
    ("PPP", 8, False,
     "¿Es de raza potencialmente peligrosa? Si es Sí, en su ficha sale el aviso de la licencia.",
     '"Sí,No"', "Sí"),
    ("Convive con niños", 13, False,
     "Sí / No / Consultar. «Consultar» sale en la ficha como «preguntadnos».",
     '"Sí,No,Consultar"', "Consultar"),
    ("Convive con perros", 13, False, "Sí / No / Consultar.", '"Sí,No,Consultar"', "Sí"),
    ("Convive con gatos", 13, False, "Sí / No / Consultar.", '"Sí,No,Consultar"', "No"),
    ("Carácter", 34, False,
     "De 3 a 5 palabras separadas por comas: «tranquilo, cariñoso, le encanta la pelota». "
     "Salen como etiquetas en su ficha.", None, "tranquilo, cariñoso, obediente"),
    ("Frase", 34, False,
     "Una frase corta en primera persona, como si la dijera él. Ej: «me flipa la pelota "
     "y dormir la siesta encima de ti».", None, "me flipa la pelota y dormir encima de ti"),
    ("Historia", 70, False,
     "De 3 a 6 líneas contando de dónde viene y cómo es. Escribid como habláis vosotras.",
     None,
     "Woody llegó en marzo de 2021 con año y medio y desde entonces no ha salido del "
     "refugio. Es un pedazo de pan: se pone panza arriba en cuanto te ve."),
    ("Fotos", 26, False,
     "Cómo se llaman las fotos que nos mandáis (ej. «woody1.jpg, woody2.jpg»). "
     "Las fotos NO se pegan aquí: se mandan aparte. La primera es la principal.",
     None, "woody1.jpg, woody2.jpg"),
    ("Notas", 42, False,
     "Cualquier cosa que nos queráis contar de esta ficha (o «ya no está» si el animal "
     "ya no debe salir en la web).", None,
     "FILA DE EJEMPLO: escribid encima o borradla"),
    ("código (no tocar)", 16, False,
     "Es el identificador interno de la ficha en la web. Lo rellenamos nosotros; en los "
     "animales nuevos se deja vacío.", None, None),
]

CENTRADAS = (0, 4, 5, 6, 8, 9, 10, 11, 12)

# ---------------------------------------------------------------- estilos

fina = Side(style="thin", color="DDD6C8")
BORDE = Border(left=fina, right=fina, top=fina, bottom=fina)


def cabecera(celda, obligatoria, ayuda):
    celda.fill = PatternFill("solid", fgColor=DORADO if obligatoria else VINO)
    celda.font = Font(name="Calibri", size=11, bold=True,
                      color=CHOCO if obligatoria else CREMA)
    celda.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    celda.border = BORDE
    if ayuda:
        nota = Comment(ayuda, "Love Animals BCN")
        nota.width, nota.height = 320, 130
        celda.comment = nota


# ---------------------------------------------------------------- hoja de animales


def hoja_animales(wb):
    ws = wb.create_sheet("Animales")
    ws.sheet_properties.tabColor = VINO

    for i, (titulo, ancho, obligatoria, ayuda, _, _) in enumerate(COLUMNAS):
        ws.column_dimensions[get_column_letter(i + 1)].width = ancho
        cabecera(ws.cell(row=1, column=i + 1, value=titulo), obligatoria, ayuda)
    ws.row_dimensions[1].height = 44
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUMNAS)), ULTIMA)

    codigo = PatternFill("solid", fgColor="EDEAE3")
    gris = Font(name="Calibri", size=9, color=GRIS)

    # fila de ejemplo, en verde y en cursiva para que se vea que no es un animal real
    ws.row_dimensions[EJEMPLO].height = 30
    for i, columna in enumerate(COLUMNAS):
        celda = ws.cell(row=EJEMPLO, column=i + 1, value=columna[5])
        celda.fill = PatternFill("solid", fgColor=CREMA)
        celda.font = Font(name="Calibri", size=10, italic=True, color=VERDE,
                          bold=(i == 17))
        celda.alignment = Alignment(vertical="center",
                                    horizontal="center" if i in CENTRADAS else "left")
        celda.border = BORDE
        if i in (4, 5):
            celda.number_format = "DD/MM/YYYY"

    for fila in range(EJEMPLO + 1, ULTIMA + 1):
        ws.cell(row=fila, column=1, value=fila - 2).font = gris
        ws.cell(row=fila, column=1).alignment = Alignment(horizontal="center", vertical="center")
        for i in range(len(COLUMNAS)):
            celda = ws.cell(row=fila, column=i + 1)
            celda.border = BORDE
            celda.alignment = Alignment(vertical="center",
                                        horizontal="center" if i in CENTRADAS else "left")
            if i in (4, 5):
                celda.number_format = "DD/MM/YYYY"
            if i == len(COLUMNAS) - 1:
                celda.fill = codigo

    for i, columna in enumerate(COLUMNAS):
        validacion = columna[4]
        if not validacion:
            continue
        letra = get_column_letter(i + 1)
        dv = DataValidation(type="list", formula1=validacion, allow_blank=True,
                            errorStyle="warning", showErrorMessage=True,
                            error="Elige uno de la lista (o déjalo en blanco si no lo sabes).",
                            errorTitle="Valor no previsto")
        ws.add_data_validation(dv)
        dv.add("%s2:%s%d" % (letra, letra, ULTIMA))

    return ws


# ---------------------------------------------------------------- instrucciones


def hoja_instrucciones(wb):
    ws = wb.create_sheet("Cómo se rellena", 0)
    ws.sheet_properties.tabColor = DORADO
    ws.sheet_view.showGridLines = False
    for letra, ancho in zip("ABCDE", (3, 26, 13, 62, 34)):
        ws.column_dimensions[letra].width = ancho

    def texto(fila, valor, col="B", hasta="E", tam=11, negrita=False, color="20211F",
              cursiva=False, alto=None):
        ws.merge_cells("%s%d:%s%d" % (col, fila, hasta, fila))
        celda = ws["%s%d" % (col, fila)]
        celda.value = valor
        celda.font = Font(name="Calibri", size=tam, bold=negrita, color=color, italic=cursiva)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        if alto:
            ws.row_dimensions[fila].height = alto
        return celda

    ws.row_dimensions[1].height = 12
    texto(2, "Fichas de los animales — Love Animals BCN", tam=20, negrita=True, color=VINO, alto=30)
    texto(3, "para que cada peludo tenga su propia página en la web nueva",
          tam=12, cursiva=True, color=DORADO, alto=20)

    texto(5, "Qué es esto", tam=13, negrita=True, color=CHOCO, alto=22)
    texto(6,
          "En la web nueva cada animal tiene su ficha: su foto grande, su edad, con quién convive, "
          "su historia y un contador de los días que lleva esperando familia. Todo eso sale de este "
          "Excel, así que lo que escribáis aquí es lo que se verá.", alto=34)
    texto(7,
          "En la hoja «Animales» hay una fila por animal. La primera fila, la verde, es un ejemplo "
          "para que se vea cómo va: escribid encima o borradla.", alto=34)

    texto(9, "Cómo se rellena", tam=13, negrita=True, color=CHOCO, alto=22)
    pasos = [
        "Una fila por animal, sin dejar filas vacías por el medio. Si se acaban las filas, se añaden "
        "más al final (o nos lo decís y os mandamos otra hoja).",
        "Las columnas con la cabecera dorada son las imprescindibles (nombre, estado y fecha de "
        "entrada). El resto ayuda mucho, pero se puede ir completando poco a poco.",
        "Las columnas con desplegable (estado, sexo, tamaño, PPP, convive…) se rellenan eligiendo de "
        "la lista, no escribiendo a mano: si no, luego no se pintan bien los iconos de la ficha.",
        "Si de algún animal no sabéis un dato, se deja en blanco: en la web sencillamente no aparece "
        "ese dato, no queda ningún hueco raro.",
        "No hace falta hacerlo de una sentada ni entregarlo entero: con que nos mandéis la hoja con "
        "lo que llevéis, ya vamos montando fichas.",
        "La última columna, «código», la rellenamos nosotros para saber qué fila es qué ficha. "
        "Vosotras la dejáis en blanco.",
    ]
    fila = 10
    for paso in pasos:
        ws["B%d" % fila] = "•"
        ws["B%d" % fila].font = Font(color=DORADO, bold=True, size=12)
        ws["B%d" % fila].alignment = Alignment(horizontal="center", vertical="top")
        texto(fila, paso, col="C", alto=30)
        fila += 1

    fila += 1
    texto(fila, "Las fotos", tam=13, negrita=True, color=CHOCO, alto=22)
    fila += 1
    texto(fila,
          "Las fotos NO se pegan en el Excel: se mandan aparte (WhatsApp, Drive, un zip… como mejor os "
          "vaya). Lo único importante es que el nombre del archivo lleve el nombre del animal y un "
          "número: woody1.jpg, woody2.jpg, woody3.jpg. La número 1 es la que sale grande en su ficha "
          "y en los álbumes.", alto=46)
    fila += 1
    texto(fila,
          "Cuantas más, mejor (con 3 por animal la ficha ya queda preciosa). Si son horizontales o "
          "cuadradas, mejor que muy alargadas.", alto=30)

    fila += 2
    texto(fila, "Qué va en cada columna", tam=13, negrita=True, color=CHOCO, alto=22)
    fila += 1
    for i, titulo in enumerate(("Columna", "¿Hace falta?", "Qué se pone", "Ejemplo")):
        celda = ws.cell(row=fila, column=i + 2, value=titulo)
        celda.fill = PatternFill("solid", fgColor=VINO)
        celda.font = Font(bold=True, color=CREMA, size=10)
        celda.alignment = Alignment(vertical="center")
        celda.border = BORDE
    ws.row_dimensions[fila].height = 20

    fila += 1
    for titulo, _, obligatoria, ayuda, _, ejemplo in COLUMNAS:
        if titulo == "Nº":
            continue
        if isinstance(ejemplo, date):
            ejemplo = ejemplo.strftime("%d/%m/%Y")
        elif titulo == "Notas":
            ejemplo = "cualquier cosa que nos queráis contar"
        elif titulo.startswith("código"):
            ejemplo = "(lo ponemos nosotros)"
        elif titulo == "Fecha de adopción":
            ejemplo = "12/06/2026"
        ws.cell(row=fila, column=2, value=titulo).font = Font(bold=True, size=10, color=CHOCO)
        c = ws.cell(row=fila, column=3, value="imprescindible" if obligatoria else "opcional")
        c.font = Font(size=10, bold=obligatoria, color=DORADO if obligatoria else GRIS)
        ws.cell(row=fila, column=4, value=ayuda).font = Font(size=10)
        ws.cell(row=fila, column=5, value=ejemplo).font = Font(size=10, italic=True, color=VERDE)
        for col in range(2, 6):
            celda = ws.cell(row=fila, column=col)
            celda.alignment = Alignment(vertical="center", wrap_text=True)
            celda.border = BORDE
            if obligatoria:
                celda.fill = PatternFill("solid", fgColor=CREMA)
        ws.row_dimensions[fila].height = 32
        fila += 1

    fila += 1
    texto(fila, "Cualquier duda, a Alex — y gracias por el trabajazo 🐾",
          tam=11, cursiva=True, color=VINO, alto=24)
    return ws


def hoja_listas(wb):
    ws = wb.create_sheet("Listas")
    ws.sheet_properties.tabColor = NUDE
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22

    ws["A1"] = "Valores que entiende la web"
    ws["A1"].font = Font(size=14, bold=True, color=VINO)
    ws["A2"] = "Los desplegables del Excel ya los llevan puestos; esto es solo para consultar."
    ws["A2"].font = Font(size=10, italic=True, color=GRIS)

    def bloque(col, fila, titulo, valores):
        c = ws.cell(row=fila, column=col, value=titulo)
        c.font = Font(bold=True, color=CREMA, size=10)
        c.fill = PatternFill("solid", fgColor=VINO)
        c.alignment = Alignment(vertical="center")
        for i, v in enumerate(valores):
            celda = ws.cell(row=fila + 1 + i, column=col, value=v)
            celda.font = Font(size=10)
            celda.border = BORDE

    bloque(1, 4, "Álbumes", ALBUMES)
    bloque(2, 4, "Estado", ["En adopción", "En acogida", "Adoptado"])
    bloque(3, 4, "Convive con…", ["Sí", "No", "Consultar"])
    bloque(2, 9, "Sexo", ["Macho", "Hembra"])
    bloque(3, 9, "Tamaño", ["Pequeño", "Mediano", "Grande"])
    bloque(2, 13, "PPP", ["Sí", "No"])
    ws["A17"] = "Si os falta algún álbum, escribidlo en «Notas» y lo creamos."
    ws["A17"].font = Font(size=10, italic=True, color=GRIS)
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    hoja_animales(wb)
    hoja_instrucciones(wb)
    hoja_listas(wb)
    wb.active = 0
    wb.save(SALIDA)
    print("hojas:", wb.sheetnames, "· filas para rellenar:", FILAS - 1)
    print("→", SALIDA)


if __name__ == "__main__":
    main()
