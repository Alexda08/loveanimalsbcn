# -*- coding: utf-8 -*-
"""Saca los enlaces de Instagram de la columna «Notas» del Excel de Carla.

Casa cada fila con su ficha por nombre + fecha de entrada (si no, solo por
nombre, y lo avisa), normaliza las URL a https://www.instagram.com/p|reel/CODIGO/,
descarta enlaces a perfiles y deja:

  insta_posts.json   {handle: [url, ...]}  -> lo aplica aplica_posts_instagram.py
  _insta_debug.json  filas sin casar, descartes, posts compartidos…

Necesita en esta carpeta excel_nuevo.xlsx (el Excel de Carla, hoja «Animales»,
cabecera en la fila 1 y datos desde la 3) y tienda_fichas.json (vuelca_fichas.py).

    python insta_extract.py
"""
import openpyxl,json,re,unicodedata,datetime,collections,sys
sys.stdout.reconfigure(encoding='utf-8')
def norm(s):
    s=unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode().lower()
    s=re.sub(r'\b(adoptad[oa]s?|no poner(lo|la)?)\b',' ',s)
    s=re.sub(r'[^a-z0-9]+',' ',s); return ' '.join(s.split())
wb=openpyxl.load_workbook('excel_nuevo.xlsx',data_only=True);ws=wb['Animales']
H=[c.value for c in ws[1]]; iN,iF,iX=H.index('Nombre *'),H.index('Fecha de entrada *'),H.index('Notas')
d=json.load(open('tienda_fichas.json',encoding='utf-8'))
fich=[]
for a in d['animales']:
    f={x['key']:x['value'] for x in a['fields']}
    fich.append(dict(handle=a['handle'],nombre=f.get('nombre',''),n=norm(f.get('nombre','')),fecha=f.get('fecha_entrada'),status=a.get('capabilities',{}).get('publishable',{}).get('status')))
byNF=collections.defaultdict(list); byN=collections.defaultdict(list)
for f in fich: byNF[(f['n'],f['fecha'])].append(f); byN[f['n']].append(f)
def fdate(v):
    if isinstance(v,datetime.datetime): return v.date().isoformat()
    if v is None: return None
    s=str(v).strip(); m=re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})',s)
    return f'{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}' if m else s[:10]
URL=re.compile(r'(?:https?://)?(?:www\.|m\.)?instagram\.com/[^\s()\[\]<>"\',]*',re.I)
POST=re.compile(r'instagram\.com/(?:[A-Za-z0-9_.]+/)?(p|reel|reels|tv)/([^/?#\s]+)',re.I)
rows=[];unmatched=[];discard=[];out={};used=collections.Counter();matched_rows={}
for r in range(3,ws.max_row+1):
    v=[c.value for c in ws[r]]
    nom=v[iN]; 
    if nom is None and not any(v): continue
    if nom is None: unmatched.append((r,nom,None,'sin nombre'));continue
    n=norm(nom); fe=fdate(v[iF]); notas=str(v[iX] or '')
    c=byNF.get((n,fe),[]); how='nombre+fecha'
    if len(c)!=1:
        if len(c)>1: unmatched.append((r,nom,fe,f'varias fichas con mismo nombre+fecha: {[x["handle"] for x in c]}'));continue
        c=byN.get(n,[]); how='solo nombre'
        if len(c)!=1:
            unmatched.append((r,nom,fe,'ninguna ficha con ese nombre' if not c else f'nombre repetido y fecha no casa: '+', '.join(f'{x["handle"]} ({x["fecha"]})' for x in c)));continue
    f=c[0]
    urls=[]
    for u in URL.findall(notas):
        m=POST.search(u)
        if not m:
            if re.search(r'instagram\.com/?$',u.split('?')[0]) or re.match(r'.*instagram\.com/[A-Za-z0-9_.]+/?(\?.*)?$',u): discard.append((r,nom,u,'enlace a perfil/no es post'))
            else: discard.append((r,nom,u,'no es post/reel'))
            continue
        kind=m[1].lower(); code=m[2]
        kind={'reels':'reel','tv':'p'}.get(kind,kind) if kind!='tv' else 'p'
        if not re.fullmatch(r'[A-Za-z0-9_-]{5,}',code): discard.append((r,nom,u,f'código no válido: {code!r}'));continue
        nu=f'https://www.instagram.com/{kind}/{code}/'
        if nu not in urls: urls.append(nu)
    # instagram mentions without URL match? e.g. "instagram" text with code
    rows.append(dict(row=r,nombre=nom,fecha=fe,handle=f['handle'],how=how,urls=urls,notas=notas))
    used[f['handle']]+=1
    if urls:
        prev=out.setdefault(f['handle'],[])
        for u in urls:
            if u not in prev: prev.append(u)
json.dump(out,open('insta_posts.json','w',encoding='utf-8'),ensure_ascii=False,indent=2)
# stats
tot=len(rows)+len(unmatched)
dist=collections.Counter(min(len(out.get(x['handle'],[])),3) for x in rows)
dup=[h for h,c in used.items() if c>1]
post2=collections.defaultdict(list)
for h,us in out.items():
    for u in us: post2[re.search(r'/(?:p|reel)/([^/]+)/',u)[1]].append(h)
shared={k:v for k,v in post2.items() if len(v)>1}
matchedh=set(used)
pub_sin=[f for f in fich if f['status']=='ACTIVE' and f['handle'] not in out]
print('filas',tot,'casadas',len(rows),'no casadas',len(unmatched),'por solo nombre',sum(x['how']=='solo nombre' for x in rows))
print('animales casados distintos',len(matchedh),'con enlaces',len(out),'total urls',sum(map(len,out.values())))
print('dist (por fila casada)',sorted(dist.items()))
print('dist por handle',sorted(collections.Counter(min(len(out.get(h,[])),3) for h in matchedh).items()))
print('fichas misma fila dup',dup)
print('fichas tienda',len(fich),'publicadas',sum(f['status']=='ACTIVE' for f in fich),'pub sin enlace',len(pub_sin))
json.dump(dict(rows=rows,unmatched=unmatched,discard=discard,shared=shared,pub_sin=pub_sin,dup=dup,fich=fich),open('_insta_debug.json','w',encoding='utf-8'),ensure_ascii=False,indent=1,default=str)
for u in unmatched: print('NOMATCH',u)
for x in discard: print('DISCARD',x)
for k,v in shared.items(): print('SHARED',k,v)
for x in rows:
    if x['how']=='solo nombre': print('SOLONOMBRE',x['row'],x['nombre'],x['fecha'],x['handle'])
